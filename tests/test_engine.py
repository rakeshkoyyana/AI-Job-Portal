from pathlib import Path

import pytest
from docx import Document

from jobpilot.ats import ats_score, interview_prep, match_probability
from jobpilot.engine import report_markdown, tailor
from jobpilot.resume_io import load_resume

ROOT = Path(__file__).resolve().parent.parent
JD = """Senior Data Engineer
Requirements:
- 3+ years of experience with Python, SQL and Apache Spark
- Experience building ETL pipelines with Airflow on AWS
- Hands-on Kubernetes experience and strong stakeholder communication
Nice to have:
- Familiarity with dbt, Terraform
"""


def make_docx(path: Path) -> Path:
    d = Document()

    def heading(t):
        r = d.add_paragraph().add_run(t)
        r.bold = True

    heading("Jane Sample")
    d.add_paragraph("City, ST | jane@example.com")
    heading("SUMMARY")
    d.add_paragraph("Data engineer with 5 years of experience building batch data pipelines.")
    heading("SKILLS")
    for label, items in [("Languages:", " SQL, Python, PySpark"), ("Cloud:", " AWS, S3, Glue"), ("Platforms:", " Databricks, Snowflake, Airflow")]:
        p = d.add_paragraph()
        p.add_run(label).bold = True
        p.add_run(items)
    heading("EXPERIENCE")
    d.add_paragraph("Data Engineer, Example Corp | Jan 2022 - Present")
    for b in ["Built PySpark ETL jobs on Databricks that cut nightly runtime by 40%",
              "Orchestrated 25 pipelines with Airflow and added data quality checks",
              "Organized quarterly team offsites and wrote onboarding docs"]:
        d.add_paragraph(b, style="List Bullet")
    heading("EDUCATION")
    d.add_paragraph("M.S. Computer Science, Sample University | 2019")
    d.save(path)
    return path


@pytest.fixture()
def src(tmp_path):
    return make_docx(tmp_path / "resume.docx")


def texts(path):
    return [p.text for p in Document(str(path)).paragraphs]


def test_format_preserved_and_truthful(src, tmp_path):
    out = tailor(src, JD, "Senior Data Engineer", "Acme", out_path=tmp_path / "o.docx")
    new = Document(str(out.docx_path))
    # labels still bold, bullets keep their list style, headers/education untouched
    skills = [p for p in new.paragraphs if p.text.startswith(("Languages:", "Cloud:", "Platforms:"))]
    assert len(skills) == 3 and all(p.runs[0].bold for p in skills)
    assert sum(p.style.name == "List Bullet" for p in new.paragraphs) == 3
    assert "Data Engineer, Example Corp | Jan 2022 - Present" in texts(out.docx_path)
    assert "M.S. Computer Science, Sample University | 2019" in texts(out.docx_path)
    # Spark is implied by PySpark/Databricks -> added. Kubernetes is NOT on the resume -> must not appear.
    blob = "\n".join(texts(out.docx_path))
    assert "Spark" in blob and "Kubernetes" not in blob
    assert "kubernetes" in out.not_added
    # relevant bullet moved ahead of the irrelevant one
    bullets = [p.text for p in new.paragraphs if p.style.name == "List Bullet"]
    assert "offsites" in bullets[-1]
    assert out.after.score >= out.before.score


def test_confirmed_skill_is_added(src, tmp_path):
    out = tailor(src, JD, "Senior Data Engineer", "Acme", confirmed=["Kubernetes"], out_path=tmp_path / "o.docx")
    assert "Kubernetes" in "\n".join(texts(out.docx_path))
    assert "kubernetes" not in out.not_added and "Kubernetes" in out.added_skills


def test_scores_and_prep(src):
    r = load_resume(src)
    rep = ats_score(r, "Senior Data Engineer", JD, docx_path=src)
    assert 0 < rep.score <= 100 and "python" in rep.matched_required and "kubernetes" in rep.missing_required
    p = match_probability(r, "Senior Data Engineer", JD)
    assert 0 < p["probability"] < 1 and p["label"]
    prep = interview_prep(r, "Senior Data Engineer", JD)
    assert prep["likely_questions"] and prep["gap_study_plan"] and prep["story_anchors_from_your_resume"]


def test_markdown_input_renders_docx(tmp_path):
    out = tailor(ROOT / "data" / "resume.md", JD, "Senior Data Engineer", "Acme", out_path=tmp_path / "m.docx")
    assert out.docx_path.exists() and any("PDF/text" in w for w in out.warnings)
    assert "Kubernetes" not in "\n".join(texts(out.docx_path))
    assert "ATS score" in report_markdown(out, "Senior Data Engineer", "Acme")

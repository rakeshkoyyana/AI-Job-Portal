"""ATS score, match probability, interview prep and suggestions.

Everything here is a transparent heuristic, not a trained model or a copy of any vendor's ATS:
the report shows how every point was earned so you can see (and argue with) the number.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from .analyzer import estimate_years, has_term, jd_requirements, required_years, supported_skills
from .resume_io import Resume

SENIORITY = {"senior", "sr", "junior", "jr", "lead", "staff", "principal", "ii", "iii", "iv", "i", "associate", "mid", "level"}
EXPECTED_SECTIONS = {"summary": ("summary", "profile", "objective"), "skills": ("skill", "competenc"),
                     "experience": ("experience", "employment"), "education": ("education",)}


@dataclass
class ATSReport:
    score: float
    breakdown: dict[str, tuple[float, float]]            # component -> (points, max)
    matched_required: list[str] = field(default_factory=list)
    missing_required: list[str] = field(default_factory=list)
    matched_preferred: list[str] = field(default_factory=list)
    missing_preferred: list[str] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)


# ------------------------------------------------------------------ formatting
def formatting_issues(docx_path: str | Path | None) -> list[str]:
    """Things that commonly break ATS parsers, read from the original DOCX."""
    if not docx_path or Path(docx_path).suffix.lower() != ".docx":
        return []
    from docx import Document

    doc = Document(str(docx_path))
    issues = []
    if doc.tables:
        issues.append(f"{len(doc.tables)} table(s): many ATS parsers read tables out of order - prefer plain paragraphs.")
    if doc.inline_shapes:
        issues.append("Images/graphics found: ATS cannot read text inside images.")
    hdr = any(p.text.strip() for s in doc.sections for p in s.header.paragraphs)
    ftr = any(p.text.strip() for s in doc.sections for p in s.footer.paragraphs)
    if hdr or ftr:
        issues.append("Text in the page header/footer: some ATS ignore it - keep contact info in the body.")
    xml = doc.element.xml
    if "w:txbxContent" in xml:
        issues.append("Text boxes found: ATS usually skips text inside text boxes.")
    if 'w:cols w:num="2"' in xml or 'w:num="2"' in xml:
        issues.append("Multi-column layout detected: can scramble reading order.")
    return issues


# ------------------------------------------------------------------ scoring
def _title_fit(resume: Resume, title: str) -> float:
    toks = [t for t in re.findall(r"[a-z+#]+", title.lower()) if t not in SENIORITY and len(t) > 1]
    if not toks:
        return 1.0
    text = resume.text.lower()
    return sum(1 for t in toks if re.search(rf"\b{re.escape(t)}", text)) / len(toks)


def _quantified_ratio(resume: Resume) -> float:
    bullets = [l for s in resume.sections for l in s.lines if l.startswith("- ")]
    if not bullets:
        return 0.0
    return sum(1 for b in bullets if re.search(r"\d", b)) / len(bullets)


def _sections_present(resume: Resume) -> tuple[float, list[str]]:
    titles = [s.title.lower() for s in resume.sections]
    missing = [name for name, keys in EXPECTED_SECTIONS.items() if not any(any(k in t for k in keys) for t in titles)]
    return 1 - len(missing) / len(EXPECTED_SECTIONS), missing


def literal_match(text: str, term: str) -> bool:
    """ATS keyword matching is literal: 'PySpark' on the page does not count as 'Spark'."""
    return re.search(rf"(?<![\w+#.]){re.escape(term)}(?![\w+#])", text, re.I) is not None


def ats_score(resume: Resume, job_title: str, jd_text: str, docx_path: str | Path | None = None) -> ATSReport:
    req, pref = jd_requirements(jd_text)
    text = resume.text
    m_req = sorted(t for t in req if literal_match(text, t))
    x_req = sorted(req - set(m_req))
    m_pref = sorted(t for t in pref if literal_match(text, t))
    x_pref = sorted(pref - set(m_pref))

    req_pts = 45 * (len(m_req) / len(req)) if req else 30.0
    pref_pts = 10 * (len(m_pref) / len(pref)) if pref else 6.0
    title_pts = 10 * _title_fit(resume, job_title)
    sec_ratio, missing_secs = _sections_present(resume)
    sec_pts = 10 * sec_ratio
    q = _quantified_ratio(resume)
    q_pts = 10 * min(1.0, q / 0.5)
    fmt_issues = formatting_issues(docx_path)
    fmt_pts = max(0.0, 10 - 3 * len(fmt_issues))
    words = len(re.findall(r"\w+", text))
    len_pts = 5.0 if 350 <= words <= 900 else 3.0 if 250 <= words <= 1100 else 1.0

    breakdown = {
        "Required keywords": (round(req_pts, 1), 45), "Preferred keywords": (round(pref_pts, 1), 10),
        "Job-title alignment": (round(title_pts, 1), 10), "Standard sections": (round(sec_pts, 1), 10),
        "Quantified achievements": (round(q_pts, 1), 10), "ATS-safe formatting": (round(fmt_pts, 1), 10),
        "Length": (len_pts, 5),
    }
    issues = list(fmt_issues)
    if missing_secs:
        issues.append(f"Missing standard section(s): {', '.join(missing_secs)}.")
    if q < 0.5:
        issues.append(f"Only {q:.0%} of bullets contain a number; aim for at least half (scale, %, time or cost saved).")
    if words < 350:
        issues.append(f"Resume is short ({words} words).")
    if words > 1100:
        issues.append(f"Resume is long ({words} words); trim to 1-2 pages.")
    return ATSReport(round(sum(p for p, _ in breakdown.values()), 1), breakdown, m_req, x_req, m_pref, x_pref, issues)


# ------------------------------------------------------------------ match probability
def _sigmoid(x: float) -> float:
    return 1 / (1 + math.exp(-x))


def match_probability(resume: Resume, job_title: str, jd_text: str, years_have: float | None = None,
                      confirmed: set[str] | None = None, location_ok: bool = True) -> dict:
    """Estimated chance this resume gets a recruiter response for this role (heuristic logistic model)."""
    rep = ats_score(resume, job_title, jd_text)
    # Candidate-fit uses EVIDENCE (implied skills count, e.g. PySpark -> Spark) plus skills the user confirmed.
    confirmed = confirmed or set()
    req, _ = jd_requirements(jd_text)
    covered = [t for t in req if has_term(resume.text, t) or t in confirmed]
    coverage = len(covered) / len(req) if req else 0.5
    title = _title_fit(resume, job_title)
    years_have = years_have if years_have is not None else estimate_years(resume.text)
    need = required_years(jd_text)
    if need is None or years_have is None:
        exp_fit, exp_note = 0.7, "years requirement not stated or not detectable"
    else:
        exp_fit = 1.0 if years_have + 1 >= need else max(0.0, 1 - (need - years_have) / 5)
        exp_note = f"asks {need}+ yrs, resume shows about {years_have:g}"
    x = 4.0 * coverage + 1.2 * title + 1.0 * exp_fit + (0.3 if location_ok else -0.8) - 3.9
    p = max(0.02, min(0.95, _sigmoid(x)))
    label = ("Strong match" if p >= 0.60 else "Competitive" if p >= 0.35 else "Stretch" if p >= 0.15 else "Long shot")
    return {"probability": round(p, 2), "label": label, "ats_score": rep.score,
            "drivers": {"required-skill coverage": f"{coverage:.0%}", "title fit": f"{title:.0%}", "experience": exp_note}}


# ------------------------------------------------------------------ interview prep
SKILL_QUESTIONS = {
    "python": "How do you structure and test a production Python data job? What do you do about memory-heavy steps?",
    "sql": "Walk through optimizing a slow query: what do you check first, and how do you prove the fix worked?",
    "spark": "How do you diagnose a skewed or slow Spark stage, and what levers do you pull?",
    "airflow": "How do you make DAGs idempotent and handle backfills and retries?",
    "snowflake": "How do you control Snowflake cost and performance (warehouse sizing, clustering, caching)?",
    "databricks": "When would you choose Delta tables and Databricks jobs over a plain Spark cluster setup?",
    "aws": "Pick one AWS pipeline you built: services involved, failure modes, and how you secured it.",
    "dbt": "How do you structure dbt models and tests, and how do you handle slowly changing dimensions?",
    "kafka": "How do you reason about delivery guarantees, ordering and consumer lag?",
    "etl": "Describe an ETL pipeline end to end: ingestion, validation, loading, monitoring, and recovery.",
    "data modeling": "Star schema vs wide tables: how do you decide, and what would you change for this team?",
    "machine learning": "Walk through a model you shipped: data, features, evaluation, and how you monitored it.",
    "llm": "How do you evaluate an LLM feature and guard against hallucination and prompt injection?",
    "kubernetes": "How would you deploy and scale a data service on Kubernetes, and what breaks first?",
    "docker": "How do you keep images small and builds reproducible for data workloads?",
}
BEHAVIORAL = [
    "Tell me about a time a pipeline or model failed in production. What did you do in the first hour and what changed afterward?",
    "Describe a disagreement with a stakeholder about requirements or priorities and how it ended.",
    "Give an example of improving something measurable (cost, speed, quality) without being asked.",
]
ASK_THEM = [
    "What does success look like in the first 90 days, and how is it measured?",
    "What is the current biggest data-quality or reliability pain, and who owns it?",
    "How does the team decide what to build next, and how much is planned vs. interrupt-driven?",
]


def interview_prep(resume: Resume, job_title: str, jd_text: str, confirmed: set[str] | None = None) -> dict:
    confirmed = confirmed or set()
    req, _ = jd_requirements(jd_text)
    sup = supported_skills(resume.text) | confirmed
    strengths = sorted(t for t in req if has_term(resume.text, t) or t in confirmed)
    missing = sorted(t for t in req if t not in strengths)
    likely = [{"topic": t, "question": SKILL_QUESTIONS.get(t, f"Tell me about a project where you used {t} and the trade-offs you made.")}
              for t in strengths[:6]]
    gaps = [{"topic": t, "plan": f"Be honest that {t} is not on your resume. Prepare: one sentence on the closest thing you've done, "
                                 f"a 2-hour hands-on mini project, and the 3 core concepts of {t}."} for t in missing[:6]]
    anchors = [l[2:] for s in resume.sections if "experience" in s.title.lower() for l in s.lines
               if l.startswith("- ") and re.search(r"\d", l)][:3]
    return {"likely_questions": likely, "gap_study_plan": gaps, "behavioral": BEHAVIORAL,
            "story_anchors_from_your_resume": anchors, "questions_to_ask_them": ASK_THEM}


# ------------------------------------------------------------------ suggestions
def suggestions(rep: ATSReport, prob: dict) -> list[str]:
    out = []
    if rep.missing_required:
        out.append("Required skills not on your resume: " + ", ".join(rep.missing_required[:10]) +
                   ". If you genuinely have any, confirm them and they will be added; otherwise do not claim them.")
    out += rep.issues
    if prob["probability"] < 0.35:
        out.append("Low estimated match: apply only if you can add a referral or a strong cover note, and prioritise better-fit roles.")
    if rep.missing_preferred:
        out.append("Nice-to-haves you could learn quickly: " + ", ".join(rep.missing_preferred[:6]) + ".")
    if not out:
        out.append("Resume is well aligned. Spend remaining effort on a referral and interview prep.")
    return out

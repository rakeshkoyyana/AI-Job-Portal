"""Resume loading (md/txt/docx/pdf) into a simple structure, and rendering back to DOCX."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

KNOWN_HEADINGS = {
    "summary", "professional summary", "profile", "objective", "skills", "technical skills",
    "core competencies", "experience", "work experience", "professional experience",
    "education", "projects", "certifications", "awards", "publications",
}
BULLET = re.compile(r"^\s*[-•*–▪●]\s+")


@dataclass
class Section:
    title: str
    lines: list[str] = field(default_factory=list)


@dataclass
class Resume:
    name: str
    contact: str
    sections: list[Section] = field(default_factory=list)

    def to_markdown(self) -> str:
        out = [f"# {self.name}", self.contact, ""]
        for s in self.sections:
            out += [f"## {s.title}", *s.lines, ""]
        return "\n".join(out).strip() + "\n"

    @property
    def text(self) -> str:
        return self.to_markdown()


def _is_heading(line: str) -> str | None:
    s = line.strip()
    if s.startswith("#"):
        return s.lstrip("#").strip()
    bare = s.rstrip(":").strip()
    if bare.lower() in KNOWN_HEADINGS:
        return bare
    if 2 < len(bare) <= 40 and bare.isupper() and not BULLET.match(s) and any(c.isalpha() for c in bare):
        return bare.title()
    return None


def parse_resume_text(text: str) -> Resume:
    lines = [l.rstrip() for l in text.splitlines()]
    while lines and not lines[0].strip():
        lines.pop(0)
    if not lines:
        raise ValueError("Resume is empty")
    name = lines[0].strip().lstrip("#").strip()
    i, contact = 1, []
    while i < len(lines) and _is_heading(lines[i]) is None:
        if lines[i].strip():
            contact.append(lines[i].strip())
        i += 1
    sections: list[Section] = []
    for line in lines[i:]:
        h = _is_heading(line)
        if h is not None:
            sections.append(Section(h))
        elif line.strip() and sections:
            sections[-1].lines.append(BULLET.sub("- ", line) if BULLET.match(line) else line.strip())
    return Resume(name, " | ".join(contact), sections)


def load_resume(path: str | Path) -> Resume:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".md", ".txt"}:
        return parse_resume_text(p.read_text(encoding="utf-8"))
    if suffix == ".docx":
        from docx import Document

        lines = []
        for para in Document(str(p)).paragraphs:
            style = (para.style.name or "").lower()
            t = para.text.strip()
            if not t:
                continue
            if style.startswith("heading") or style == "title":
                lines.append("## " + t)
            elif "list" in style:
                lines.append("- " + t)
            else:
                lines.append(t)
        return parse_resume_text("\n".join(lines))
    if suffix == ".pdf":
        from pypdf import PdfReader

        text = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
        return parse_resume_text(text)
    raise ValueError(f"Unsupported resume format: {suffix}")


def render_docx(resume: Resume, path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.7)
        s.top_margin = s.bottom_margin = Inches(0.6)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(resume.name)
    r.bold, r.font.size = True, Pt(18)
    p = doc.add_paragraph(resume.contact)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for sec in resume.sections:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        hr = h.add_run(sec.title.upper())
        hr.bold, hr.font.size = True, Pt(11.5)
        for line in sec.lines:
            if line.startswith("- "):
                para = doc.add_paragraph(line[2:], style="List Bullet")
            else:
                para = doc.add_paragraph()
                para.add_run(line).bold = not line.startswith("-") and len(line) < 110 and sec.title.lower() not in {"summary", "professional summary", "profile"}
            para.paragraph_format.space_after = Pt(1)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def render_text_docx(text: str, path: str | Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
    for block in text.strip().split("\n\n"):
        doc.add_paragraph(block.strip())
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path

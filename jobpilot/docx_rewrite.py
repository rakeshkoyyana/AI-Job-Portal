"""Rewrite a resume IN ITS OWN FORMAT.

Works on a neutral list of paragraphs. For a .docx we edit the original file paragraph-by-paragraph
(keeping every run style, bullet, font and margin); for .md/.txt/.pdf input we rebuild the structure and
render a clean .docx. Only three kinds of paragraphs are ever edited: the summary, skills lines and
experience bullets. Role titles, employers, dates and education lines are never touched.
"""
from __future__ import annotations

import copy
import json
import re
from dataclasses import dataclass
from pathlib import Path

from .analyzer import find_skills, supported_skills
from .resume_io import BULLET, Resume, Section, _is_heading, load_resume, render_docx

EDITABLE = {"summary", "skills_line", "bullet"}

ACRONYMS = {"aws", "gcp", "sql", "etl", "elt", "nlp", "llm", "rag", "ci/cd", "iam", "s3", "emr", "dbt", "api"}
DISPLAY = {
    "pyspark": "PySpark", "postgresql": "PostgreSQL", "mongodb": "MongoDB", "dynamodb": "DynamoDB", "bigquery": "BigQuery",
    "github actions": "GitHub Actions", "power bi": "Power BI", "tensorflow": "TensorFlow", "pytorch": "PyTorch",
    "fastapi": "FastAPI", "graphql": "GraphQL", "rest api": "REST API", "a/b testing": "A/B Testing", "scikit-learn": "scikit-learn",
    "hugging face": "Hugging Face", "mlops": "MLOps", "sql server": "SQL Server", "delta lake": "Delta Lake",
    "data warehousing": "Data Warehousing", "data modeling": "Data Modeling", "data quality": "Data Quality",
    "machine learning": "Machine Learning", "deep learning": "Deep Learning", "data pipelines": "Data Pipelines",
}
CATEGORY_HINTS = [  # (label keywords, skills that belong there)
    (("language", "programming"), {"python", "sql", "scala", "java", "javascript", "typescript", "golang", "rust", "bash", "shell scripting"}),
    (("cloud", "aws"), {"aws", "gcp", "azure", "s3", "glue", "lambda", "iam", "redshift", "athena", "emr", "kinesis", "sagemaker", "bigquery"}),
    (("platform", "tool", "framework", "technolog", "big data", "database", "data"), {"databricks", "snowflake", "airflow", "dbt", "kafka", "spark", "flink", "dagster", "prefect", "tableau", "power bi", "looker"}),
    (("practice", "method", "concept", "other"), {"etl", "elt", "data modeling", "data quality", "ci/cd", "git", "orchestration", "data governance", "streaming", "data pipelines"}),
]


def display(term: str) -> str:
    if term in DISPLAY:
        return DISPLAY[term]
    if term in ACRONYMS:
        return term.upper()
    return term.title() if " " in term else term.capitalize()


@dataclass
class Para:
    idx: int
    text: str
    section: str          # lowercased section title ("" before the first heading)
    kind: str             # heading | header | summary | skills_line | bullet | other


# ------------------------------------------------------------------ building Para lists
def _kind_for(section: str, text: str, is_bullet: bool) -> str:
    s = section.lower()
    if any(k in s for k in ("summary", "profile", "objective")):
        return "summary"
    if "skill" in s or "competenc" in s:
        return "skills_line"
    if any(k in s for k in ("experience", "employment", "project")):
        return "bullet" if is_bullet else "header"
    return "other"


def paras_from_docx(path: str | Path) -> list[Para]:
    from docx import Document

    out, section = [], ""
    for i, p in enumerate(Document(str(path)).paragraphs):
        t = p.text.strip()
        if not t:
            out.append(Para(i, "", section, "other"))
            continue
        style = (p.style.name or "").lower()
        h = _is_heading(t) if not BULLET.match(t) else None
        if style.startswith("heading") or h:
            section = (h or t).lower().rstrip(":")
            out.append(Para(i, t, section, "heading"))
            continue
        ppr = p._p.pPr
        is_bullet = "list" in style or bool(BULLET.match(t)) or (ppr is not None and ppr.numPr is not None)
        out.append(Para(i, t, section, _kind_for(section, t, is_bullet) if section else "other"))
    return out


def paras_from_resume(resume: Resume) -> list[Para]:
    out = [Para(0, resume.name, "", "other"), Para(1, resume.contact, "", "other")]
    for s in resume.sections:
        out.append(Para(len(out), s.title, s.title.lower(), "heading"))
        for line in s.lines:
            is_b = line.startswith("- ")
            out.append(Para(len(out), line[2:] if is_b else line, s.title.lower(), _kind_for(s.title, line, is_b)))
    return out


def resume_from_paras(base: Resume, paras: list[Para], edits: dict[int, str]) -> Resume:
    sections, cur = [], None
    for p in paras[2:]:
        if p.kind == "heading":
            cur = Section(p.text)
            sections.append(cur)
        elif cur is not None and (p.text or p.idx in edits):
            text = edits.get(p.idx, p.text)
            cur.lines.append(("- " + text) if p.kind == "bullet" else text)
    return Resume(base.name, base.contact, sections)


# ------------------------------------------------------------------ planning (rules)
def _jd_relevance(text: str, jd_terms: set[str]) -> int:
    return len(supported_skills(text) & jd_terms)


def _split_skill_line(text: str) -> tuple[str, list[str]]:
    head, sep, rest = text.partition(":")
    if not sep:
        return "", [x.strip() for x in text.split(",") if x.strip()]
    return head.strip(), [x.strip() for x in rest.split(",") if x.strip()]


def _join_skill_line(head: str, items: list[str]) -> str:
    return f"{head}: {', '.join(items)}" if head else ", ".join(items)


def _best_line(lines: list[tuple[int, str, list[str]]], term: str) -> int:
    for keys, members in CATEGORY_HINTS:
        if term in members:
            for idx, head, _ in lines:
                if any(k in head.lower() for k in keys):
                    return idx
    return lines[-1][0]


def plan_rule_edits(paras: list[Para], jd_terms: set[str], matched: list[str], add_terms: list[str]) -> dict[int, str]:
    edits: dict[int, str] = {}
    # skills: reorder + add supported/confirmed terms
    skill_lines = [(p.idx, *_split_skill_line(p.text)) for p in paras if p.kind == "skills_line"]
    pending = {t: None for t in add_terms}
    new_items: dict[int, list[str]] = {}
    for idx, head, items in skill_lines:
        new_items[idx] = list(items)
    if skill_lines:
        for t in add_terms:
            tgt = _best_line([(i, h, it) for i, h, it in skill_lines], t)
            if not any(display(t).lower() == x.lower() for x in new_items[tgt]):
                new_items[tgt].append(display(t))
    for idx, head, _ in skill_lines:
        items = new_items[idx]
        items.sort(key=lambda it: 0 if find_skills(it) & jd_terms or it.lower() in jd_terms else 1)
        new = _join_skill_line(head, items)
        if new != next(p.text for p in paras if p.idx == idx):
            edits[idx] = new
    # summary: add a truthful strengths sentence
    for p in paras:
        if p.kind == "summary" and matched and "Relevant strengths for this role" not in p.text:
            top = ", ".join(display(t) for t in matched[:6])
            edits[p.idx] = p.text.rstrip() + f" Relevant strengths for this role: {top}."
            break
    # bullets: within each consecutive group, put the most JD-relevant text into the earlier slots
    group: list[Para] = []

    def flush():
        if len(group) > 1:
            texts = sorted((g.text for g in group), key=lambda t: -_jd_relevance(t, jd_terms))
            for g, t in zip(group, texts):
                if g.text != t:
                    edits[g.idx] = t
        group.clear()

    for p in paras:
        if p.kind == "bullet":
            group.append(p)
        else:
            flush()
    flush()
    return edits


# ------------------------------------------------------------------ guardrail for any edit set
def verify_edits(paras: list[Para], edits: dict[int, str], allowed_skills: set[str]) -> list[str]:
    by_idx = {p.idx: p for p in paras}
    issues = []
    orig_numbers = set(re.findall(r"\d[\d,.]*%?", " ".join(p.text for p in paras)))
    for idx, new in edits.items():
        p = by_idx.get(idx)
        if p is None:
            issues.append(f"edit targets unknown paragraph {idx}")
            continue
        if p.kind not in EDITABLE:
            issues.append(f"edit to protected paragraph ({p.kind}): {p.text[:40]!r}")
            continue
        extra = find_skills(new) - allowed_skills
        if extra:
            issues.append(f"unsupported skills {sorted(extra)} in: {new[:50]!r}")
        nums = set(re.findall(r"\d[\d,.]*%?", new)) - orig_numbers
        if nums:
            issues.append(f"new numbers {sorted(nums)} in: {new[:50]!r}")
    return issues


# ------------------------------------------------------------------ LLM planning (optional)
LLM_SYSTEM = """You tailor a resume to a job description by editing individual paragraphs.
HARD RULES:
- You are given editable paragraphs as JSON [{id, kind, text}]. Return ONLY JSON: {"edits": {"<id>": "<new text>"}} for paragraphs you change.
- Never add a skill, tool, number, metric, employer, title or date that is not already in the original text. You may only use the extra skills listed in ALLOWED_EXTRA.
- You may reword bullets and the summary to use the job's vocabulary where it truthfully describes the same work, and may reorder items in skills lines.
- Keep bullet length similar to the original; keep skills lines in the same 'Label: a, b, c' form."""


def plan_llm_edits(client, model: str, paras: list[Para], jd_text: str, title: str, allowed_extra: list[str],
                   matched: list[str]) -> dict[int, str]:
    editable = [{"id": p.idx, "kind": p.kind, "text": p.text} for p in paras if p.kind in EDITABLE]
    user = (f"JOB TITLE: {title}\nJOB DESCRIPTION:\n{jd_text[:6000]}\n\nSKILLS TO EMPHASISE (already on resume): {matched}\n"
            f"ALLOWED_EXTRA: {allowed_extra}\n\nEDITABLE PARAGRAPHS:\n{json.dumps(editable)}")
    msg = client.messages.create(model=model, max_tokens=4000, system=LLM_SYSTEM, messages=[{"role": "user", "content": user}])
    raw = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text").strip()
    raw = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.M).strip()
    data = json.loads(raw)["edits"]
    return {int(k): str(v).strip() for k, v in data.items() if str(v).strip()}


# ------------------------------------------------------------------ applying edits
def _remove_runs(runs):
    for r in runs:
        r._element.getparent().remove(r._element)


def set_paragraph_text(p, new_text: str) -> None:
    """Replace a paragraph's text but keep the formatting of its runs (label runs in 'Label: items' stay bold, etc.)."""
    runs = list(p.runs)
    if not runs:
        p.add_run(new_text)
        return
    old = p.text
    if ":" in old and ":" in new_text and old.split(":", 1)[0] == new_text.split(":", 1)[0]:
        label_len = old.index(":") + 1
        body = new_text[label_len:]
        cum = 0
        for i, r in enumerate(runs):
            if cum + len(r.text) >= label_len:
                k = label_len - cum
                if k == len(r.text) and i + 1 < len(runs):
                    runs[i + 1].text = body
                    _remove_runs(runs[i + 2:])
                else:
                    r.text = r.text[:k] + body
                    _remove_runs(runs[i + 1:])
                return
            cum += len(r.text)
    runs[0].text = new_text
    _remove_runs(runs[1:])


def apply_docx_edits(src: str | Path, dst: str | Path, edits: dict[int, str]) -> Path:
    from docx import Document

    doc = Document(str(src))
    pars = doc.paragraphs
    for idx, text in edits.items():
        set_paragraph_text(pars[idx], text)
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return Path(dst)


def docx_has_table_text(path: str | Path) -> bool:
    from docx import Document

    return any(c.text.strip() for t in Document(str(path)).tables for r in t.rows for c in r.cells)
